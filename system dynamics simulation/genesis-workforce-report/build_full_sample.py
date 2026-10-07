"""Build the full Genesis report sample .docx, following Raj's v2 outline,
with the user's corrections baked in."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import csv

G = "/home/hatch/workspace/genesis/"
d = Document()
st = d.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(11)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def bold(cell):
    for p in cell.paragraphs:
        for r in p.runs: r.bold = True

def P(text, style="Normal"):
    return d.add_paragraph(text, style=style)

def H(text, lvl=1):
    return d.add_heading(text, level=lvl)

def B(text):
    d.add_paragraph(text, style="List Bullet")

def mktable(headers, rows, highlight_row=None):
    tb = d.add_table(rows=1 + len(rows), cols=len(headers))
    tb.style = "Light Grid Accent 1"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = tb.rows[0].cells[j]; c.text = h; bold(c)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            tb.rows[i + 1].cells[j].text = str(v)
        if highlight_row == i + 1:
            for c in tb.rows[i + 1].cells: shade(c, "D9EAD3")
    return tb

# ---------------- Title ----------------
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Genesis Workforce and Production Report"); r.bold = True; r.font.size = Pt(20)
s = d.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
s.add_run("ISYE 681 \u2013 Introduction to System Dynamics and Applications").font.size = Pt(13)
a = d.add_paragraph(); a.alignment = WD_ALIGN_PARAGRAPH.CENTER
a.add_run("Farnood Tavakoli, Ibrahim Oyeyinka, Raj Nooti").font.size = Pt(12)

# ---------------- 1. Problem Statement ----------------
H("1. Problem Statement")
P("Genesis leadership set two 10-year targets: double the workforce from 850 to 1,700 employees, "
  "and triple annual production from 2.6 million to 7.8 million lamps.")
P("Every new hire enters as a Rookie (1,000 lamps/year). With experience, workers mature into "
  "Experienced staff (2,000 lamps/year) and eventually Gray Hairs (4,000 lamps/year). Reaching Gray "
  "Hair status takes seven years under baseline timing: two years as a Rookie plus five as Experienced.")
P("The targets pull in different directions. Tripling output while only doubling headcount means average "
  "output per employee must rise from 3,059 to 4,588 lamps/year \u2014 a 50% increase \u2014 even as rapid "
  "hiring floods the pipeline with Rookies, the least productive tier. The question is whether any "
  "workforce policy can deliver both targets at once.")

# ---------------- 2. Model and baseline ----------------
H("2. System Dynamics Model and Baseline Steady State")
P("We modeled the workforce in Vensim as an aging chain of three stocks \u2014 Rookies (R), "
  "Experienced (E), and Gray Hairs (G) \u2014 connected by hiring, learning, maturing, and retirement flows:")
P("Hiring Rate = 50 + STEP(50, 1)  (50/year in year 1, 100/year thereafter)", style="List Bullet")
P("Learning Rate = Rookies / Learning Time", style="List Bullet")
P("Maturing Rate = Experienced / Maturing Time", style="List Bullet")
P("Retiring Rate = Gray Hairs / Retiring Time", style="List Bullet")
P("Production = 1,000\u00b7Rookies + 2,000\u00b7Experienced + 4,000\u00b7Gray Hairs  (lamps/year)", style="List Bullet")
d.add_picture(G + "report/figures/fig1_schematic.png", width=Inches(5.5))
P("Figure 1 \u2013 Aging-chain stock-and-flow structure.", style="Normal").italic = True
H("Initial conditions and baseline steady state", lvl=2)
P("The case study starts the system in equilibrium: Rookies = 100 (learning time 2 yrs), Experienced = 250 "
  "(maturing time 5 yrs), Gray Hairs = 500 (retiring time 10 yrs). At these values the learning rate "
  "(100/2), maturing rate (250/5), and retirement rate (500/10) are each 50 employees/year, exactly "
  "balancing the 50/year hiring rate \u2014 so the 850-person workforce and 2.6M-lamp output are a true "
  "steady state.")
P("Note on baseline convention: this report uses the case-study steady state (learning 2 yrs, maturing 5 yrs, "
  "retiring 10 yrs) as the baseline. The .mdl file\u2019s default learning time is 1 yr; all results below "
  "use the stated (2, 5, 10) baseline unless a factor level says otherwise.",
  style="Normal").italic = True
H("Baseline 10-year simulation", lvl=2)
P("Simulating the baseline policy (hiring steps from 50 to 100/year in year 2; baseline time constants) "
  "over ten years reaches only 1,233 employees and 3.72 million lamps/year \u2014 well short of both "
  "management targets.")
d.add_picture(G + "report/figures/fig2_baseline.png", width=Inches(5.8))
P("Figure 2 \u2013 Baseline trajectory: workforce and production over 10 years.", style="Normal").italic = True

# ---------------- 3. Feasibility ----------------
H("3. Management Targets and Feasibility Analysis")
P("For the management target, Genesis must produce 7.8 million lamps with 1,700 employees: "
  "4,588 lamps/person/year, versus 3,059 today \u2014 an implied 50% productivity improvement "
  "((4588 \u2212 3059)/3059).")
H("Is this goal achievable?", lvl=2)
P("No \u2014 not through hiring and workforce development alone. Two structural constraints bind:")
B("Output cap by employee type. Productivity is capped at 1,000 / 2,000 / 4,000 lamps/year across the "
  "three tiers, and the target needs 4,588 per person \u2014 above even the Gray-Hair maximum. "
  "Even 1,700 all-Gray-Hair workers could produce only 6.8M lamps, so the 7.8M target is mathematically "
  "infeasible at current productivity.")
B("Workforce dilution. Gray Hairs cannot be hired directly; every worker enters as a Rookie and needs "
  "seven years to reach Gray-Hair status. Doubling the workforce therefore floods the pipeline with "
  "Rookies, dragging average productivity down exactly when it must rise.")
P("Because of these two constraints, no hiring profile or timing policy can deliver 7.8M lamps with "
  "1,700 employees unless productivity itself grows. The experimental design below quantifies how far "
  "timing policies can go \u2014 and how large the remaining productivity gap is.")

# ---------------- 4. Experimental design (user's part) ----------------
H("4. Experimental Design")
P("To evaluate how internal delays affect workforce growth and output, we tested an 8-scenario "
  "(2\u00b3 factorial) design in Vensim. The three factors are:")
B("Learning time (x1): 1 year (low) vs. 2 years (high)")
B("Maturing time (x2): 3 years (low) vs. 5 years (high)")
B("Retiring time (x3): 10 years (low) vs. 15 years (high)")
P("Combination 7 (2, 5, 10) is the case-study baseline. In the design runs hiring is not stepped as in "
  "the baseline; instead the post-year-1 hiring rate is solved per combination to hit a stated target, "
  "as described under each table.")
H("Table 1 \u2013 Hiring solved for 7.8M lamps/year", lvl=2)
P("Each combination\u2019s hiring rate was back-solved so the model produces exactly 7.8M lamps/year at "
  "current productivity. Production is therefore fixed by construction in every row \u2014 the meaningful "
  "response is the workforce each combination needs:")
rows1 = list(csv.reader(open(G + "genesis_8combo_corrected.csv")))
data1 = [[r[0], r[1], r[2], r[3], r[4], r[5], r[6]] for r in rows1[1:]]
mktable(["Combo", "x1 (yr)", "x2 (yr)", "x3 (yr)", "Hiring/yr after yr 1",
         "Employees yr 10", "Production yr 10 (M)"], data1, highlight_row=2)
d.add_picture(G + "report/figures/fig3_combos.png", width=Inches(5.8))
P("Figure 3 \u2013 Workforce required per combination to reach 7.8M lamps (combination 2 highlighted).",
  style="Normal").italic = True
H("Table 2 \u2013 Hiring solved for 1,700 employees", lvl=2)
P("Here each combination\u2019s hiring rate was back-solved to land exactly 1,700 employees at year 10. "
  "The table shows the resulting base-production and the productivity gain still required to reach 7.8M:")
rows2 = list(csv.reader(open(G + "genesis_tableB_gain.csv")))
data2 = [[r[0], r[1], r[2], r[3], r[4], r[5], r[6] + "%"] for r in rows2[1:]]
mktable(["Combo", "x1 (yr)", "x2 (yr)", "x3 (yr)", "Hiring/yr for 1,700",
         "Base prod. (M)", "Gain needed"], data2, highlight_row=2)
H("Analysis of results", lvl=2)
B("Combination 2 (1, 3, 15) is the best timing configuration on both criteria: it needs the fewest "
  "employees to hit 7.8M at base productivity (2,482), and the smallest productivity gain (+41.7%) "
  "under the 1,700 headcount. It accelerates onboarding to 1 year, shortens maturing to 3 years, "
  "and extends Gray-Hair tenure to 15 years.")
B("Retiring time (x3) has the highest leverage of the three factors. Extending Gray-Hair tenure from "
  "10 to 15 years adds about 210K lamps/year on average across the design (at fixed 1,700 headcount), "
  "because it keeps the most productive tier \u2014 4,000 lamps/year \u2014 in the workforce longer.")
B("The structural gap persists under every timing combination: even the best case peaks at 5.51M lamps "
  "at base productivity. Timing policy alone cannot reach 7.8M with 1,700 employees \u2014 a productivity "
  "intervention is mathematically necessary.")
d.add_picture(G + "report/figures/fig4_recommended.png", width=Inches(5.8))
P("Figure 4 \u2013 Recommended policy (combination 2, ~147 hires/yr): workforce and production trajectories.",
  style="Normal").italic = True

# ---------------- 5. Trade-offs ----------------
H("5. Trade-offs and Potential Alternatives")
P("The simulations reveal a hard trade-off: Genesis cannot satisfy both goals with hiring and timing "
  "policies alone. Leadership must choose what to prioritize \u2014 or change the production function itself.")
H("Prioritize the 7.8M production goal", lvl=2)
P("Hitting 7.8M lamps at current productivity requires 2,482\u20133,031 employees depending on timing "
  "(Table 1) \u2014 far above the 1,700 cap. Even the best timing (combination 2) needs 2,482 heads and "
  "244.6 hires/year, roughly 2.5\u00d7 the baseline hiring rate.")
H("Prioritize the 1,700 headcount", lvl=2)
P("Capping headcount at 1,700 means holding hiring to about 147 Rookies/year under combination 2 timing. "
  "Output then reaches only 5.51M lamps at base productivity \u2014 a 2.29M shortfall that must be closed "
  "by raising output per worker.")
H("Potential alternatives", lvl=2)
B("Accommodate higher headcount: drop the 1,700 cap and staff to ~2,480 under combination-2 timing. "
  "This hits the production target but substantially raises labor cost and dilutes average productivity.")
B("Invest in automation and tools: add the required 41.7% productivity boost. Combination 2 with 1,700 "
  "employees yields 5.51M at base productivity; the 41.7% gain closes the gap to 7.8M. The mechanisms "
  "(technology, working methods, automation, cycle-time improvement) and their costs are outside this model.")
B("Lower the production target: recalibrate to what timing policy alone can deliver \u2014 about 5.5M "
  "lamps with 1,700 employees under the best timing.")

# ---------------- 6. Recommendation ----------------
H("6. Recommendation and Conclusion")
P("We recommend the following package:")
B("Adopt combination 2 timing: 1-year learning, 3-year maturing, 15-year Gray-Hair tenure \u2014 the "
  "highest-leverage timing policy in the design.")
B("Hire ~147 Rookies/year after year 1 to reach exactly 1,700 employees by year 10.")
B("Invest in automation and working methods to lift productivity 42% by year 10, closing the "
  "5.51M \u2192 7.8M gap.")
B("Decouple the headcount and production targets in planning: the analysis proves they cannot both be "
  "met on hiring and timing alone, so future plans should treat productivity growth as an explicit lever.")
P("In short: the best timing policy gets Genesis to 1,700 people and 5.51M lamps; a 42% productivity "
  "program buys the rest. Without it, management must choose between fewer lamps and more heads.")

# ---------------- Appendix ----------------
H("Appendix \u2013 Vensim Model Equations")
P("Stocks (initial values): Rookies = 100, Experienced = 250, Gray Hairs = 500.", style="List Bullet")
P("Hiring Rate = 50 + STEP(50, 1)", style="List Bullet")
P("Learning Rate = Rookies / Learning Time", style="List Bullet")
P("Maturing Rate = Experienced / Maturing Time", style="List Bullet")
P("Retiring Rate = Gray Hairs / Retiring Time", style="List Bullet")
P("Production = 1000\u00b7Rookies + 2000\u00b7Experienced + 4000\u00b7Gray Hairs", style="List Bullet")
P("Baseline: learning 2 yrs, maturing 5 yrs, retiring 10 yrs. Simulation horizon 10 years, "
  "time step 0.0625 year.", style="List Bullet")

d.save(G + "genesis_full_sample.docx")
print("saved genesis_full_sample.docx")
